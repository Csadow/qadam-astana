"""Create matching, editable PPTX + vector PDF and timed Russian speaker notes."""
from pathlib import Path
import json, textwrap
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from PIL import Image

ROOT=Path(__file__).resolve().parents[1]; OUT=ROOT/'presentation'; OUT.mkdir(exist_ok=True)
pdfmetrics.registerFont(TTFont('Sans','/usr/share/fonts/noto/NotoSans-Regular.ttf'))
pdfmetrics.registerFont(TTFont('Bold','/usr/share/fonts/noto/NotoSans-Bold.ttf'))
pdfmetrics.registerFont(TTFont('Serif','/usr/share/fonts/noto/NotoSerif-Regular.ttf'))
INK='#203e36'; PAPER='#f6f5ef'; LIME='#d6ed92'; MUTED='#687768'; SOFT='#e8eddd'; WHITE='#fffefa'; AMBER='#91652d'
S=.6; W=1600; H=900
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
pdf=canvas.Canvas(str(OUT/'Qadam_ArcSolo_pitch_RU.pdf'),pagesize=(W*S,H*S))
pdf.setTitle('Qadam — ArcSolo | Build with AI, Track 1');pdf.setAuthor('ArcSolo')
notes=[]; slide=None; page=0

def rgb(color):return RGBColor.from_string(color.lstrip('#'))
def rect(x,y,w,h,fill,line=None,radius=False):
    shape=slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,Pt(x*S),Pt(y*S),Pt(w*S),Pt(h*S))
    shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill)
    if radius: shape.adjustments[0]=.04
    for effect in shape._element.xpath('.//a:effectRef'): effect.set('idx','0')
    shape.line.fill.background() if not line else None
    if line:shape.line.color.rgb=rgb(line)
    pdf.setFillColor(fill);pdf.setStrokeColor(line or fill)
    if radius:pdf.roundRect(x*S,(H-y-h)*S,w*S,h*S,10*S,fill=1,stroke=bool(line))
    else:pdf.rect(x*S,(H-y-h)*S,w*S,h*S,fill=1,stroke=bool(line))
    return shape

def text(value,x,y,w,size=30,color=INK,font='Sans',lineheight=1.28):
    lines=[]
    for paragraph in str(value).split('\n'):
        if not paragraph:lines.append('');continue
        line=''
        for word in paragraph.split(' '):
            test=(line+' '+word).strip()
            if pdfmetrics.stringWidth(test,font,size*S)>w*S and line:lines.append(line);line=word
            else:line=test
        lines.append(line)
    box=slide.shapes.add_textbox(Pt(x*S),Pt(y*S),Pt(w*S),Pt((len(lines)*size*lineheight+12)*S))
    tf=box.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Noto Serif' if font=='Serif' else 'Noto Sans';p.font.size=Pt(size*S);p.font.bold=font=='Bold';p.font.color.rgb=rgb(color);p.space_after=Pt(0);p.space_before=Pt(0);p.line_spacing=Pt(size*lineheight*S)
        pdf.setFillColor(color);pdf.setFont(font,size*S);pdf.drawString(x*S,(H-y-size*.98-i*size*lineheight)*S,line)
    return y+len(lines)*size*lineheight

def image(path,x,y,w,h):
    iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);nw,nh=iw*scale,ih*scale
    slide.shapes.add_picture(str(path),Pt(x*S),Pt(y*S),width=Pt(nw*S),height=Pt(nh*S))
    pdf.drawImage(str(path),x*S,(H-y-nh)*S,nw*S,nh*S,mask='auto')

def line(x1,y1,x2,y2,color=INK,width=6):
    shape=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1*S),Pt(y1*S),Pt(x2*S),Pt(y2*S))
    shape.line.color.rgb=rgb(color);shape.line.width=Pt(width*S)
    pdf.setStrokeColor(color);pdf.setLineWidth(width*S);pdf.line(x1*S,(H-y1)*S,x2*S,(H-y2)*S)

def new(title,kicker,dark=False):
    global slide,page
    if page:pdf.showPage()
    page+=1;slide=prs.slides.add_slide(prs.slide_layouts[6]);rect(0,0,W,H,INK if dark else PAPER)
    color=LIME if dark else MUTED
    text('QADAM  /  ARCSOLO',70,42,600,18,color,'Bold')
    text(kicker.upper(),1040,42,490,17,color)
    text(title,70,115,1450,55,LIME if dark else INK,'Serif',1.15)
    text('BUILD WITH AI  ·  SMART CITY  ·  TRACK 1',70,846,1100,16,color)
    text(f'{page:02d}',1460,842,65,21,color,'Bold')

def note(title,seconds,script):
    notes.append({'slide':page,'title':title,'seconds':seconds,'script':script})
    slide.notes_slide.notes_text_frame.text=f'{seconds} секунд\n\n{script}'

new('Кафе недоступно. Вернуться нужно к 15:00.','01 / Идея',True)
text('qadam.',75,285,950,150,LIME,'Serif')
text('Планы меняются.\nСрок возвращения остаётся.',85,495,1150,49,WHITE,'Serif')
text('ИИ-помощник по городским прогулкам',85,683,1000,27,'#b6c6b7')
rect(1210,298,220,220,LIME,radius=True)
line(1262,465,1375,352);line(1290,352,1375,352);line(1375,352,1375,437)
note('Знакомство',25,'Здравствуйте! Мы — ArcSolo, это Qadam. Представьте: у вас час возле EXPO, но кафе оказалось недоступно. Вы потеряли ещё пятнадцать минут, а вернуться нужно к трём. Qadam перестраивает оставшуюся прогулку, сохраняя первоначальный срок и точку возвращения. Покажем это в репетиции. Маршрут опирается на реальные места, источники и явные ограничения; время остаётся оценкой.')

new('Свободное время есть. Готового плана — нет.','02 / Проблема')
for x,n,title,body in [(70,'01','Что рядом?','Список мест ещё не учитывает вашу точку старта и интересы.'),(565,'02','Успею ли я?','На дорогу, посещение и возвращение нужен общий бюджет.'),(1060,'03','Можно ли доверять?','Розетки, часы работы и доступность часто неизвестны.')]:
    rect(x,300,465,340,WHITE,radius=True);text(n,x+28,330,390,28,MUTED,'Serif');text(title,x+28,401,390,35,INK,'Serif');text(body,x+28,474,390,27,MUTED)
text('Для жителей, родителей и гостей города — когда есть 30–120 минут.',76,715,1400,30,INK)
note('Проблема',25,'Недостаточно назвать популярный парк. Пользователь спрашивает: что подходит рядом со мной и под мой запас времени? Родителю важен формат прогулки, туристу — расстояние, человеку с ноутбуком — розетка. Но данных о розетке может не быть. Мы соединяем поиск, проверку условий и короткое объяснение, а неизвестное оставляем видимым.')

new('Один экран: от желания к маленькой прогулке.','03 / Продукт')
image(OUT/'assets/map.png',70,270,940,380)
image(OUT/'assets/route.png',70,655,940,145)
for y,a,b in [(285,'Контекст','Старт, время, настроение, компания и ограничения.'),(450,'1–3 варианта','Факты, личная идея и маленькая миссия открытия.'),(615,'План с источниками','Групповой маршрут, план Б и проверяемые источники.')]:
    text(a,1060,y,465,32,INK,'Serif');text(b,1060,y+53,445,26,MUTED)
note('Продукт',30,'Это рабочий прототип. Задаём старт, время и желание. Получаем один маршрут: карточки, карта и экспорт показывают те же остановки и длительность. На телефоне маршрут виден первым, подробные предпочтения раскрываются по необходимости. Демо использует явно обозначенное время; режим «Сейчас» берёт текущие часы. Интерфейс работает на английском, русском и казахском. Геолокация необязательна.')

new('Демо: 60 минут возле EXPO.','04 / Показать продукт')
rect(70,275,1460,120,INK,radius=True);text('53 минуты = 33 минуты пешком + 20 минут на месте',105,307,1370,39,LIME,'Serif')
image(OUT/'assets/map.png',70,435,850,340)
text('Короткая пауза → Аллея мира и согласия → старт',975,450,540,37,INK,'Serif')
text('Возвращение уже включено.\nОстановки сокращены до 10 минут.\nЗакрепить · заменить · удалить · отменить.',975,575,525,27,MUTED)
text('EXPO: сплошные линии по OSM; пунктир — оценка соединений. Входы и перекрытия не проверены.',77,798,1430,21,MUTED)
note('Демонстрация',60,'Покажем час возле EXPO. Две остановки занимают примерно пятьдесят три минуты, включая возвращение. Сплошная линия следует связным пешеходным путям OSM, пунктир — примерным соединениям. На первом экране нажимаем «Попробовать Идём вместе»: репетиция уже настроена на возвращение к пятнадцати часам с пятью минутами запаса. Представим, что пришли в кафе, а оно недоступно. Отмечаем это, затем теряем пятнадцать минут. Точка изменилась, остаток времени уменьшился, но срок пятнадцать часов сохранён. Добавляем дождь: уличные посещения исчезают, ограничения остаются. Можно выбрать возвращение сейчас. Указание недоступности действует только в нашей прогулке, общую базу оно не меняет. Реальное состояние входов требует проверки.')

new('Данные собраны автоматически. Пробелы измерены.','05 / Данные')
for x,num,label in [(70,'300','именованных POI'),(565,'32','факта из источников'),(1060,'19','мест с обогащением')]:
    rect(x,290,465,245,WHITE,radius=True);text(num,x+28,315,400,90,INK,'Serif');text(label,x+30,460,400,26,MUTED)
text('OSM автоматически + отдельные факты официальных источников',76,592,1450,35,INK,'Serif')
text('5 районов · 102 записи с часами · 26 тегов доступности. 32 факта о 18 местах.\nИсточники уточнили расписания и исключили места с ограниченным доступом.',76,666,1430,27,MUTED)
text('Пять районов: 5 188 пешеходных путей · Снимок 18.10.2026 · © OpenStreetMap contributors / ODbL.',76,783,1450,21,MUTED)
note('Данные',30,'Сбор данных автоматизирован. Скрипт получает OpenStreetMap, фильтрует категории и удаляет дубликаты. В снимке триста мест в пяти районах. У ста двух есть часы, у двадцати шести — тег доступности. Отдельный слой обогащает девятнадцать записей: тридцать два факта о восемнадцати местах и ограничения доступа. Первичные источники уточнили понедельник в музее и вход в кампус. Дата исследования не означает проверку на месте.')

new('RAG: сначала найти и проверить, потом предложить.','06 / Подход')
items=[('01','POI','Реальный снимок\nи источники'),('02','Поиск','BM25 + понятия\nна трёх языках'),('03','Ограничения','Дорога, время,\nчасы и доступность'),('04','Gemini','Объяснение, идея\nи маленькая миссия'),('05','Проверка','ID мест и цитат,\nформат и ограничения')]
for i,(n,title,body) in enumerate(items):
    x=70+i*295;rect(x,300,270,270,INK if i==3 else WHITE,radius=True);color=LIME if i==3 else INK;text(n,x+23,327,220,24,color);text(title,x+23,389,230,31,color,'Serif');text(body,x+23,462,230,22,'#c4d1c0' if i==3 else MUTED)
text('Живой Gemini — локально и в видео; публичное демо без платной генерации.',75,625,1440,29,INK)
text('Факты, маршруты и часы храним отдельно от сгенерированных советов.\nПроверка ссылок не доказывает истинность каждой фразы. При ошибке — шаблоны.',75,697,1440,26,MUTED)
note('RAG',35,'Поиск сочетает BM25 и многоязычные понятия. Проверки учитывают расстояние, длительность, расписание и отказ от категорий. Gemini получает только выбранные места, факты и ограниченный контекст. Он пишет объяснение, занятие и небольшую миссию. Проверяем ID мест, цитаты и распространённые неподтверждённые обещания. Но это не доказательство каждого предложения. Факты и правила показаны отдельно, а при ошибке остаётся обозначенный шаблонный режим.')

new('Кафе недоступно. Потеряли 15 минут. Что дальше?','07 / Идём вместе')
image(OUT/'assets/companion.png',70,265,750,540)
text('Тот же срок. Новая точка.',845,290,665,39,INK,'Serif')
text('Остаток маршрута пересчитан от подтверждённого места. Возвращение — к исходному старту.',845,355,655,29,MUTED)
text('Запас не исчезает молча',845,495,665,35,INK,'Serif')
text('Если уже не успеваем, видны опоздание и нехватка резерва. Можно сразу возвращаться.',845,555,655,28,MUTED)
text('Для жюри — явная репетиция',845,695,665,31,INK,'Serif')
text('Для прогулки — часы устройства и ручное подтверждение. Без фонового GPS.',845,745,650,24,MUTED)
note('Адаптивная прогулка',35,'Главный новый сценарий: мы пришли в кафе, а оно недоступно. Затем потеряли пятнадцать минут и начался дождь. Qadam пересчитывает только оставшийся план от подтверждённой точки. Срок возвращения и первоначальный старт не меняются, запас не тратится молча. Завершённые места не добавляются снова. Если посещение уже не помещается, остаётся возвращение; если и оно не успевает, приложение прямо показывает это. Жюри может проверить всё в явно обозначенной репетиции. В реальной прогулке часы обновляются при действиях, местоположение подтверждает человек.')

new('Доверие — это видимое «не знаем».','08 / Этика и приватность')
image(OUT/'assets/coffee.png',70,280,840,440)
text('Факты и идеи',975,285,520,38,INK,'Serif');text('Источник POI отделён от предложения, что сделать.',975,348,510,27,MUTED)
text('Локация без истории',975,465,520,38,INK,'Serif');text('Расчёт в браузере. Точные координаты и исходный запрос не отправляются в Gemini.',975,530,510,27,MUTED)
text('Никаких обещаний о тишине, розетках, ночной безопасности или маршруте без ступеней без данных.',75,762,1440,26,AMBER)
note('Доверие',30,'Отсутствие данных не превращается в уверенное обещание. Розетки и тишина здесь неизвестны. Тег доступности заведения не означает проверенный путь без ступеней. Точные координаты, исходный запрос и история редактирования остаются в браузере. Карта не обращается к сторонним тайлам. В Gemini уходят публичные места и нормализованные предпочтения. Сетевые метаданные при этом могут обрабатываться инфраструктурой провайдера.')

new('Проверили поведение, а не только красивый экран.','09 / Проверка')
for x,y,num,label in [(70,275,'103','тестов движка, API и Worker'),(815,275,'160','комбинаций ограничений'),(70,520,'25 / 25','приёмочных сценариев'),(815,520,'106 / 106','проверок в браузере')]:
    rect(x,y,715,215,WHITE,radius=True);text(num,x+30,y+20,640,79,INK,'Serif');text(label,x+34,y+142,640,27,MUTED)
text('Ещё 6 тестов сбора графа и 100 сочетаний согласованности маршрута. Пользовательский пилот впереди.',76,778,1440,23,MUTED)
note('Проверка',25,'Проходят сто три теста кода, шесть проверок сборщика графа, двадцать пять приёмочных сценариев и сто шесть проверок браузера. Целевой аудит сложных запросов улучшился с трёх до десяти из десяти; ещё двенадцать перефразировок прошли. Проверены обходы, односторонние пути, разрывы сети, дедлайны и неполные истории. Живой Gemini работает. Это инженерная проверка конкретных примеров, не общая точность и не пользовательское исследование.')

new('Сначала польза человеку. Затем городской сервис.','10 / Бизнес-гипотеза')
rect(70,280,710,450,INK,radius=True);text('Бесплатно для прогулки',105,320,635,43,LIME,'Serif');text('Жители, родители, туристы.\nКороткий план без регистрации.',105,423,625,31,WHITE);text('Ценность: меньше времени на выбор,\nбольше понятных действий рядом.',105,570,625,27,'#c1cebc')
rect(815,280,715,450,WHITE,radius=True);text('B2B после пилота',850,320,640,43,INK,'Serif');text('Отели, городские порталы,\nорганизаторы мероприятий.',850,423,625,31,MUTED);text('Гипотеза: платный виджет и интеграция.\nРеклама маркируется и не отменяет ограничения.',850,565,625,27,MUTED)
text('План пилота: 20 пользователей · 3 сценария · 2 недели. Это цель, не уже достигнутый результат.',76,779,1440,23,MUTED)
note('Бизнес',30,'Пользовательская прогулка остаётся бесплатной. B2B-гипотеза — виджет для отелей, городских порталов и организаторов мероприятий. Выручку мы пока не заявляем. Для пилота предлагаем двадцать пользователей, три сценария и две недели. Измеряем время до выбора, полезность и ошибки ограничений. Партнёрские размещения должны быть обозначены и не смогут отменять ограничения маршрута.')

new('Честные ограничения превращаем в план развития.','11 / Следующий шаг')
rows=[('Данные неполные','Проверенные описания и расписания от партнёров.'),('Входы не проверены','Пять районов на карте; проверить входы и перекрытия.'),('В тексте возможны ошибки','Оценка смысла ответов и полезности с пользователями.'),('Язык остаётся эвристикой','Больше примеров и проверка казахского языка.')]
for i,(a,b) in enumerate(rows):
    y=285+i*117;rect(70,y,1460,93,WHITE,radius=True);text(a,100,y+23,590,29,INK,'Serif');text('→',708,y+20,70,34,MUTED);text(b,795,y+26,700,26,MUTED)
text('Понимаем «полчаса» и 1h30m. Неясное время просим подтвердить.',76,787,1430,24,MUTED)
note('Развитие',30,'Основные возможности работают: точнее понимаем время и отказ, строим пешеходные пути в пяти районах, связываем места в истории, предлагаем проверенный выход из пустого результата и учитываем срок возвращения с запасом. Теперь пути собраны в пяти районах, а не только в EXPO. Добавлены местные факты о пяти местах, неясное время требует подтверждения, карточки показывают реальный переход от предыдущей остановки. Входы не проверены на месте; нужны обновление источников и проверка казахского. Пользовательские сессии отложены; спрос и удобство пока не доказаны.')

new('Один район. Один час. Один хороший Qadam.','12 / Финал',True)
text('Не просто «куда сходить».\nПлан, которому понятны границы.',80,290,1400,69,LIME,'Serif',1.2)
text('Личные идеи. Истории города. Пешеходные пути.\nСопровождение прогулки и возвращение с запасом.',85,535,1320,35,WHITE)
text('ArcSolo  ·  github.com/Csadow/qadam-astana',85,717,1400,29,'#c1cebc')
text('Код и планировщик доступны публично. Живой Gemini — в локальном приложении и видео.',85,779,1400,21,'#c1cebc')
note('Финал',20,'Qadam превращает свободный час в небольшой понятный план. Пользователь видит реальные места, ограничения и возможность изменить решение. Работают сбор данных, интерфейс, генерация и проверки. План адаптируется уже во время прогулки: меняется оставшийся путь, сохраняется исходное обязательство вернуться. Жюри может проверить это в репетиции. Спасибо! Готовы показать любой сценарий.')

prs.save(OUT/'Qadam_ArcSolo_pitch_RU.pptx');pdf.save()
elapsed=0;lines=['# Qadam · ArcSolo — Russian pitch notes','',f'Target rehearsal duration: {sum(n["seconds"] for n in notes)//60}:{sum(n["seconds"] for n in notes)%60:02d}. Includes a short live demo. Adjust to your speaking speed.','', 'Live Gemini is demonstrated in the local app and current recording. Recheck the model badge before presenting. The repository and judge preview are public; the preview uses retrieval/templates without paid generation and links to the recording. Organizer submission and YouTube upload have not been performed, at the participant’s request. Participant testing is deferred.','']
for n in notes:
    end=elapsed+n['seconds'];lines.extend([f'## {n["slide"]:02d}. {n["title"]} · {elapsed//60}:{elapsed%60:02d}–{end//60}:{end%60:02d}','',n['script'],'']);elapsed=end
(OUT/'SPEAKER_NOTES_RU.md').write_text('\n'.join(lines),encoding='utf-8')
(OUT/'slides.json').write_text(json.dumps(notes,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(f'Created {page} editable slides + PDF; target {elapsed} seconds.')
